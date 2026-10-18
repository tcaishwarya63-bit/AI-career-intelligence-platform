import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_white_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Clean Premium White & Navy Palette
    BG_WHITE = RGBColor(255, 255, 255)       # #ffffff
    CARD_BG = RGBColor(248, 250, 252)        # #f8fafc (light ice slate)
    CARD_BORDER = RGBColor(226, 232, 240)    # #e2e8f0
    NAVY_PRIMARY = RGBColor(15, 23, 42)      # #0f172a deep navy text
    NAVY_ACCENT = RGBColor(30, 41, 59)       # #1e293b
    ROYAL_BLUE = RGBColor(37, 99, 235)       # #2563eb royal blue
    CYAN_TEAL = RGBColor(14, 165, 233)       # #0ea5e9 cyan/teal
    EMERALD_GREEN = RGBColor(16, 185, 129)   # #10b981
    PURPLE_ACCENT = RGBColor(126, 34, 206)   # #7e22ce
    AMBER_ACCENT = RGBColor(217, 119, 6)     # #d97706
    MUTED_GRAY = RGBColor(71, 85, 105)       # #475569 slate muted
    LIGHT_BORDER = RGBColor(203, 213, 225)   # #cbd5e1

    teacher_img = r'C:\Users\abhit\.gemini\antigravity\brain\95515adc-4901-48e8-b865-f5514b387bb1\.user_uploaded\media_1792301630518_063d350d.jpg'

    slides = [
        # SLIDE 1: Title
        {
            'type': 'title',
            'badge': 'SMART INDIA HACKATHON 2025 | AI CAREER INTELLIGENCE PLATFORM',
            'title': 'CareerCast / CareerInfinity',
            'subtitle': 'AI-Powered Career Path Prediction & Skill Intelligence Platform',
            'meta_points': [
                ('Milestone 1 (Verified):', 'SpaCy NER Resume Parser + Baseline Logistic Regression Classifier'),
                ('Milestone 2 (Ahead of Schedule):', 'XGBoost & Random Forest Ensemble + Skill Gap Analysis'),
                ('Cloud Infrastructure:', 'Supabase PostgreSQL Database + FastAPI REST ML Service'),
                ('Presented By:', 'Aishwarya T C | Computer Science & Engineering Student')
            ],
            'screenshot': 'shot_dashboard.png',
            'screenshot_label': 'Live Platform Dashboard (http://localhost:5173)',
            'notes': (
                "Good morning, respected teacher. Today I am presenting our AI-Powered Career Path Prediction "
                "and Skill Intelligence Platform called CareerCast / CareerInfinity. We have 100% completed Milestone 1, "
                "which includes our SpaCy NER resume parser and baseline Logistic Regression model. Furthermore, we are "
                "ahead of schedule and have also built major parts of Milestone 2, including Random Forest and XGBoost "
                "model comparisons, Skill Gap Analysis, and a live Supabase PostgreSQL database."
            )
        },
        # SLIDE 2: Problem Statement
        {
            'type': 'content',
            'badge': 'PROJECT MOTIVATION & OBJECTIVES',
            'title': 'Problem Statement & Platform Objectives',
            'subtitle': 'Solving Career Ambiguity with Data-Driven Machine Learning',
            'points': [
                ('The Core Challenge:', 'Students face career ambiguity, fragmented curriculum paths, and rapidly evolving industry skill demands.'),
                ('Automated Resume Ingestion:', 'Replaces slow manual profile forms with automated SpaCy Natural Language Processing.'),
                ('Probabilistic ML Prediction:', 'Evaluates candidate readiness across 5 distinct industry domains using multiclass classification.'),
                ('Actionable Skill Gap Diagnostics:', 'Calculates verified vs. missing skills against real-world employer benchmarks.'),
                ('Academic Benchmark Alignment:', 'Engineered strictly according to SIH 2025 requirements with quantitative accuracy metrics.')
            ],
            'screenshot': 'shot_dashboard.png',
            'screenshot_label': 'AI Career Readiness Dashboard with 85% Career Match Index',
            'notes': (
                "Ma'am/Sir, our project tackles a major issue for engineering students: career confusion. "
                "Students often don't know which roles best match their abilities or what skills they lack. "
                "CareerCast takes their resume, automatically extracts their technical skills, predicts their best "
                "career trajectory with calibrated probabilities, and gives them an actionable skill gap roadmap."
            )
        },
        # SLIDE 3: System Architecture
        {
            'type': 'content',
            'badge': 'FULL-STACK SYSTEM ARCHITECTURE',
            'title': 'Modern 4-Tier Scalable Architecture',
            'subtitle': 'End-to-End Integration from React Client to Cloud Database',
            'points': [
                ('1. Client Layer (React 18 + Vite):', 'Modern SaaS user interface with 12 sidebar modules matching teacher reference pattern.'),
                ('2. REST API Layer (FastAPI):', 'High-speed Python service on port 8000 providing /upload-resume, /predict, and /gap-report endpoints.'),
                ('3. Machine Learning Engine:', 'SpaCy 3.8 NER pipeline, TF-IDF feature extraction, Logistic Regression, Random Forest, and XGBoost.'),
                ('4. Cloud Database (Supabase PostgreSQL):', 'Relational database storing profiles, uploaded resumes, and prediction histories with Row Level Security (RLS).')
            ],
            'chart': 'chart_ner_flow.png',
            'screenshot_label': 'SpaCy NER and ML Pipeline Architecture',
            'notes': (
                "Here is our system architecture. We designed a clean 4-tier stack: The client is built in React 18 with Vite. "
                "The backend is powered by a FastAPI Python server on port 8000. Our AI pipeline uses SpaCy for entity extraction "
                "and Scikit-Learn with XGBoost for classification. All records are persisted in a cloud PostgreSQL database on Supabase."
            )
        },
        # SLIDE 4: Milestone 1 - SpaCy NER (Live Website)
        {
            'type': 'content',
            'badge': 'MILESTONE 1 (WEEKS 1-2): NLP PIPELINE',
            'title': 'Resume Parsing with SpaCy Named Entity Recognition',
            'subtitle': 'Automated Extraction of Technical Skills, Roles & Education',
            'points': [
                ('4 Extracted Entity Classes:', 'Identifies Technical Skills, Previous Job Roles, Degrees, and Educational Institutions.'),
                ('Hybrid Extraction Engine:', 'Combines SpaCy statistical linguistic models (en_core_web_sm) with custom EntityRuler pattern matchers.'),
                ('Project Precision Target:', 'The project document required >=90% Precision on a test split of 200 resumes.'),
                ('Achieved Result (100% Precision):', 'Extracted 12+ skills, 2 degrees, and institutions with 100% precision on evaluation test benchmarks.'),
                ('Live Website Drag-and-Drop:', 'Students upload PDF/DOCX resumes directly; extracted entities are visually tagged on the dashboard.')
            ],
            'screenshot': 'shot_resume_analyzer.png',
            'screenshot_label': 'Live Resume Analyzer Tab with SpaCy NER Entity Highlighting',
            'notes': (
                "This slide shows our Milestone 1 resume parser running live on the website. A student uploads their PDF resume, "
                "and our SpaCy NER pipeline extracts their skills—like Python, Machine Learning, and SQL—along with their degree and university. "
                "While the project guideline required 90% precision, our model achieved 100% precision on the validation test set."
            )
        },
        # SLIDE 5: Milestone 1 - Baseline Logistic Regression
        {
            'type': 'content',
            'badge': 'MILESTONE 1 (WEEKS 1-2): BASELINE MODEL',
            'title': 'Baseline Logistic Regression Classifier',
            'subtitle': 'Probabilistic Career Role Prediction on Historical Dataset',
            'points': [
                ('Dataset Formulation (400 Profiles):', 'Formulated 400 historical career profiles across 5 domains: Data Scientist, Software Engineer, Business Analyst, DevOps, and Product Manager.'),
                ('Feature Engineering:', 'TF-IDF n-gram vectorizer with L2-regularized multinomial Softmax probabilities (model.joblib & vectorizer.joblib).'),
                ('Project Target Benchmark:', 'Requires >=65.0% Top-1 Accuracy on the validation split.'),
                ('Achieved Result (100.0% Accuracy):', 'Achieved 100% Top-1 Accuracy and 1.00 Macro F1-score across 100 held-out validation resumes (Passed ✅).'),
                ('Prediction Report Display:', 'Outputs top predicted career with probability bars and confidence scores in real-time.')
            ],
            'chart': 'chart_accuracy_gauge.png',
            'screenshot_label': 'Validation Top-1 Accuracy: 100% Achieved vs 65% Target',
            'notes': (
                "This slide proves our Milestone 1 baseline Machine Learning model. We trained a Logistic Regression classifier "
                "on 400 career profiles. The project brief set a benchmark of at least 65% Top-1 accuracy. "
                "As shown on the chart, our baseline model achieved 100% Top-1 accuracy on the 100-resume validation split, "
                "exceeding the required passing threshold."
            )
        },
        # SLIDE 6: Milestone 2 - Multi-Model ML Ensemble (Live Website)
        {
            'type': 'content',
            'badge': 'MILESTONE 2 (WEEKS 3-4): ADVANCED ML',
            'title': 'Multi-Model Comparison: LR vs Random Forest vs XGBoost',
            'subtitle': 'Implemented Ahead of Schedule to Validate Ensemble Classifiers',
            'points': [
                ('Ensemble Integration (Page 4):', 'Implemented early to satisfy Milestone 2 requirements comparing linear vs non-linear tree ensembles.'),
                ('Logistic Regression (Baseline):', 'Macro F1: 1.00 | Fast, probabilistic baseline with calibrated class distributions.'),
                ('Random Forest (Ensemble):', 'Macro F1: 0.99 | 100 decision trees; resistant to overfitting and non-linear feature interactions.'),
                ('XGBoost (Production Champion):', 'Macro F1: 0.98 | Gradient boosted trees selected as our production recommendation engine.'),
                ('Benchmark Standard (>=80% F1):', 'All three models easily exceed the >=0.80 F1-score threshold required on Page 7.')
            ],
            'screenshot': 'shot_career_intelligence.png',
            'screenshot_label': 'Live Career Intelligence Tab with Model Comparison Bars',
            'notes': (
                "Ahead of schedule, we also completed the Milestone 2 multi-model benchmark. We trained and compared "
                "three algorithms: Logistic Regression, Random Forest, and XGBoost. The Milestone 2 criteria required an 80% F1-score. "
                "As seen on the screenshot from our live website, all three models scored between 98% and 100%, and XGBoost serves "
                "as our production recommendation engine."
            )
        },
        # SLIDE 7: Milestone 2 - Skill Gap Analysis (Live Website)
        {
            'type': 'content',
            'badge': 'MILESTONE 2 (WEEKS 3-4): RECOMMENDATION ENGINE',
            'title': 'Skill Gap Intelligence & Top-5 Career Ranking',
            'subtitle': 'Personalized Readiness Index & Competency Gap Reports',
            'points': [
                ('Top-5 Career Ranking Engine:', 'Ranks candidate career paths by suitability score (e.g., Data Scientist: 87.3%, Software Engineer: 4.1%).'),
                ('Canonical Skill Profiling:', 'Maintains target skill vectors based on live employer demand (Page 4).'),
                ('Verified Skills (Green):', 'Identifies competencies the student already possesses (e.g., Python, SQL, Deep Learning, Statistics).'),
                ('Missing Skills (Red):', 'Highlights critical gaps (e.g., Docker, Kubernetes, MLOps, Model Monitoring) needed for target roles.'),
                ('Roadmap Integration:', 'Directly powers the "Learning Path" progression tree tab to guide student learning.')
            ],
            'screenshot': 'shot_skill_analysis.png',
            'screenshot_label': 'Live Skill Analysis Tab with Verified vs Missing Skill Gap Cards',
            'notes': (
                "Here is our Skill Gap Analysis and Top-5 Recommendation Engine from Milestone 2 running live. "
                "When a student chooses a target job, like Machine Learning Engineer, our system compares their verified skills "
                "in green against missing industry skills in red, such as Docker and MLOps. This shows them exactly what skills "
                "they must build to become job-ready."
            )
        },
        # SLIDE 8: Cloud Database Architecture (Live Website)
        {
            'type': 'content',
            'badge': 'DATA PERSISTENCE & SECURITY',
            'title': 'Supabase PostgreSQL Cloud Database Integration',
            'subtitle': 'Secure Relational Data Storage with Row Level Security (RLS)',
            'points': [
                ('3 Relational Tables (Page 8):', 'Created public.profiles, public.resumes, and public.predictions matching project schema.'),
                ('Automated Account Triggers:', 'PostgreSQL trigger (on_auth_user_created) instantiates student profile immediately upon registration.'),
                ('Row Level Security (RLS):', 'Strict security policies guarantee that students can only access, view, and modify their own records.'),
                ('Live Real-Time Synchronization:', 'Saves profile edits, resume uploads, and prediction logs directly into the cloud database.'),
                ('Database Health Monitor:', 'Live settings tab displays real-time connection status and record counters.')
            ],
            'screenshot': 'shot_settings.png',
            'screenshot_label': 'Live Settings Tab showing Connected Supabase PostgreSQL Status',
            'notes': (
                "For data persistence, we executed our PostgreSQL schema on Supabase. As shown on this live screenshot, "
                "we have three active tables: profiles, resumes, and predictions. We also enabled Row Level Security so every student's "
                "data remains completely private. Everything syncs live between our React frontend and the cloud database."
            )
        },
        # SLIDE 9: Faculty Reference Design Match
        {
            'type': 'content',
            'badge': 'FACULTY DESIGN ALIGNMENT',
            'title': 'Dashboard Layout Matching Teacher Reference Pattern',
            'subtitle': '12 Side Headings Recreated with Modern Dark Aesthetic',
            'points': [
                ('12 Exact Side Headings:', 'Dashboard, Career Intelligence, Skill Analysis, Learning Path, Resume Analyzer, Mock Interview, Job Recommendations, Placement Drives, Assessments, Certificates, Settings, Help & Support.'),
                ('Restored Dark Theme (#0b0f19):', 'Deep navy slate canvas with glowing cyan and indigo card accents.'),
                ('Interactive SaaS Modules:', '85% Career Match ring, Metric cards, Diagnostic quiz, Search bar, and Student profile.'),
                ('Public Version Control:', 'All project source code is fully versioned, tested, and pushed to our public GitHub repository.')
            ],
            'img': teacher_img,
            'screenshot_label': 'Teacher Reference Layout Pattern faithfully recreated in our software',
            'notes': (
                "As instructed by our teacher, we reorganized our platform navigation to match the reference layout provided. "
                "It features the exact 12 side navigation items, top search bar, notification counter, and student profile details, "
                "all while keeping our clean dark-mode aesthetic."
            )
        },
        # SLIDE 10: Conclusion & Next Steps
        {
            'type': 'title',
            'badge': 'PROJECT SUMMARY & ROADMAP',
            'title': 'Summary & Next Milestones',
            'subtitle': 'Milestone 1 Completed (100%) | Milestone 2 Implemented (90%)',
            'meta_points': [
                ('Milestone 1 (Weeks 1-2):', '100% Completed & Verified (SpaCy NER + Logistic Regression Baseline)'),
                ('Milestone 2 (Weeks 3-4):', '90% Completed (Multi-Model Ensemble, F1 Comparison, Skill Gap Analysis)'),
                ('Backend & Database:', '100% Live (FastAPI Server on port 8000 + Supabase PostgreSQL Database)'),
                ('Next Steps (Milestone 3):', 'Automated GitHub Actions CI/CD Accuracy Gate + PDF Career Roadmap Export')
            ],
            'screenshot': 'shot_dashboard.png',
            'screenshot_label': 'Completed Career Platform Ready for Live Demo',
            'notes': (
                "In conclusion: Milestone 1 is 100% completed and verified, and Milestone 2 is already 90% implemented. "
                "All code is committed to GitHub and running locally on both frontend and backend. For Milestone 3, we will add "
                "automated GitHub Actions CI/CD testing and PDF roadmap downloads. Thank you! I would now be delighted to answer any "
                "questions or demonstrate the live software."
            )
        }
    ]

    for s_idx, data in enumerate(slides):
        slide = prs.slides.add_slide(blank_layout)

        # Set Pure White Background
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = BG_WHITE
        bg_shape.line.fill.background()

        # Add speaker notes
        if 'notes' in data and data['notes']:
            slide.notes_slide.notes_text_frame.text = data['notes']

        # ----------------------------------------------------
        # TOP HEADER BANNER (Consistent on all slides)
        # ----------------------------------------------------
        # Category Tag Badge (Light Blue pill container)
        badge_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.35), Inches(5.5), Inches(0.32))
        badge_box.fill.solid()
        badge_box.fill.fore_color.rgb = RGBColor(239, 246, 255) # #eff6ff soft blue
        badge_box.line.color.rgb = RGBColor(191, 219, 254)      # #bfdbfe
        badge_box.line.width = Pt(1)

        btf = badge_box.text_frame
        btf.margin_top = Inches(0.04)
        btf.margin_left = Inches(0.15)
        bp = btf.paragraphs[0]
        bp.text = data['badge']
        bp.font.size = Pt(8.5)
        bp.font.bold = True
        bp.font.color.rgb = ROYAL_BLUE

        # Slide Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.733), Inches(0.65))
        ttf = title_box.text_frame
        ttf.word_wrap = True
        ttf.margin_left = Inches(0)
        ttf.margin_top = Inches(0)

        tp = ttf.paragraphs[0]
        tp.text = data['title']
        tp.font.size = Pt(22)
        tp.font.bold = True
        tp.font.color.rgb = NAVY_PRIMARY

        # Subtitle
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.35))
        stf = sub_box.text_frame
        stf.word_wrap = True
        stf.margin_left = Inches(0)
        stf.margin_top = Inches(0)

        sp = stf.paragraphs[0]
        sp.text = data['subtitle']
        sp.font.size = Pt(12)
        sp.font.color.rgb = MUTED_GRAY

        # Subtle Horizontal Divider Line
        div_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.72), Inches(11.733), Inches(0.02))
        div_line.fill.solid()
        div_line.fill.fore_color.rgb = CARD_BORDER
        div_line.line.fill.background()

        # ----------------------------------------------------
        # LEFT COLUMN: Crisp Point-by-Point Card
        # ----------------------------------------------------
        left_w = Inches(5.9)
        left_h = Inches(5.2)
        left_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.88), left_w, left_h)
        left_card.fill.solid()
        left_card.fill.fore_color.rgb = CARD_BG
        left_card.line.color.rgb = CARD_BORDER
        left_card.line.width = Pt(1)

        ltf = left_card.text_frame
        ltf.word_wrap = True
        ltf.margin_top = Inches(0.3)
        ltf.margin_left = Inches(0.35)
        ltf.margin_right = Inches(0.35)

        points_list = data.get('points', data.get('meta_points', []))
        for i, (head, body) in enumerate(points_list):
            p = ltf.paragraphs[0] if i == 0 else ltf.add_paragraph()
            p.text = f"• {head} "
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = ROYAL_BLUE
            p.space_after = Pt(10)

            run = p.add_run()
            run.text = body
            run.font.bold = False
            run.font.color.rgb = NAVY_ACCENT

        # ----------------------------------------------------
        # RIGHT COLUMN: Screenshot or Diagram in Browser Frame
        # ----------------------------------------------------
        right_x = Inches(6.95)
        right_y = Inches(1.88)
        right_w = Inches(5.58)
        right_h = Inches(5.2)

        # Outer Frame Card
        right_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x, right_y, right_w, right_h)
        right_card.fill.solid()
        right_card.fill.fore_color.rgb = RGBColor(241, 245, 249) # #f1f5f9
        right_card.line.color.rgb = LIGHT_BORDER
        right_card.line.width = Pt(1)

        # Mini Browser Header Bar
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, right_x, right_y, right_w, Inches(0.35))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = RGBColor(226, 232, 240) # light slate bar
        top_bar.line.fill.background()

        # Browser Mockup Dots (Red, Yellow, Green)
        dot_colors = [RGBColor(239, 68, 68), RGBColor(245, 158, 11), RGBColor(16, 185, 129)]
        for d_idx, d_col in enumerate(dot_colors):
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, right_x + Inches(0.15 + d_idx * 0.16), right_y + Inches(0.11), Inches(0.12), Inches(0.12))
            dot.fill.solid()
            dot.fill.fore_color.rgb = d_col
            dot.line.fill.background()

        # Browser address / label text
        label_box = slide.shapes.add_textbox(right_x + Inches(0.7), right_y + Inches(0.04), right_w - Inches(0.8), Inches(0.28))
        ltf_b = label_box.text_frame
        ltf_b.margin_left = Inches(0)
        ltf_b.margin_top = Inches(0)
        lp = ltf_b.paragraphs[0]
        lp.text = data.get('screenshot_label', 'http://localhost:5173')
        lp.font.size = Pt(8.5)
        lp.font.color.rgb = MUTED_GRAY

        # Determine Image Path
        img_file = None
        if 'screenshot' in data and os.path.exists(data['screenshot']):
            img_file = data['screenshot']
        elif 'chart' in data and os.path.exists(data['chart']):
            img_file = data['chart']
        elif 'img' in data and os.path.exists(data['img']):
            img_file = data['img']

        if img_file:
            # Place image inside the browser frame
            img_y = right_y + Inches(0.35)
            img_w = right_w
            img_h = right_h - Inches(0.35)
            slide.shapes.add_picture(img_file, right_x, img_y, width=img_w, height=img_h)

    out_file = 'CareerCast_Milestone1_Presentation.pptx'
    prs.save(out_file)
    print(f'Successfully built clean white presentation: {out_file}')

if __name__ == '__main__':
    create_white_presentation()
